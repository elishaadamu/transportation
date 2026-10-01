#!/usr/bin/env python3
import os
import shutil
import subprocess
import docx
from docx.text.paragraph import Paragraph

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGINAL_DOCX = os.path.join(BASE_DIR, "items to use/TCAMPO_Plan2050_Ch1-Ch2_20260916-zak.docx")
PUBLIC_DOWNLOADS = os.path.join(BASE_DIR, "public/downloads")
SECTIONS_DIR = os.path.join(PUBLIC_DOWNLOADS, "sections")
TMP_DIR = "/tmp/tcampo_conversion"

os.makedirs(PUBLIC_DOWNLOADS, exist_ok=True)
os.makedirs(SECTIONS_DIR, exist_ok=True)
os.makedirs(TMP_DIR, exist_ok=True)

CHROME_PATH = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

def docx_to_pdf_via_ql(docx_path, output_pdf_path):
    preview_dir = docx_path + ".qlpreview"
    if os.path.exists(preview_dir):
        shutil.rmtree(preview_dir)
    
    # Run QuickLook manage
    subprocess.run(["qlmanage", "-p", "-o", os.path.dirname(docx_path), docx_path], capture_output=True)
    preview_html = os.path.join(preview_dir, "Preview.html")
    
    if not os.path.exists(preview_html):
        raise RuntimeError(f"QuickLook failed to produce Preview.html for {docx_path}")
    
    # Print to PDF using Chrome headless without URL header/footer
    cmd = [
        CHROME_PATH,
        "--headless",
        "--disable-gpu",
        "--run-all-compositor-stages-before-draw",
        "--no-pdf-header-footer",
        f"--print-to-pdf={output_pdf_path}",
        f"file://{preview_html}"
    ]
    subprocess.run(cmd, capture_output=True)
    if not os.path.exists(output_pdf_path) or os.path.getsize(output_pdf_path) == 0:
        raise RuntimeError(f"Chrome failed to create PDF at {output_pdf_path}")
    
    print(f"  [OK] Generated {os.path.basename(output_pdf_path)} ({os.path.getsize(output_pdf_path):,} bytes)")

def cleanup_old_ai_docs():
    print("1. Cleaning up any old AI-generated or temporary files...")
    old_files = [
        "TCAMPO_Plan2050_Chapter3_Transportation_System_Today.pdf",
        "TCAMPO_Plan2050_Chapter4_System_Tomorrow_2050.pdf",
        "TCAMPO_Plan2050_Chapter5_Funding_the_Plan.pdf",
        "TCAMPO_Plan2050_Appendices.pdf",
        "TCAMPO_Plan2050_Ch1-Ch2.docx",
    ]
    for folder in [PUBLIC_DOWNLOADS, os.path.join(BASE_DIR, "dist/downloads")]:
        for f in old_files:
            p = os.path.join(folder, f)
            if os.path.exists(p):
                os.remove(p)
                print(f"  Removed {p}")

    # Remove outdated test section PDFs if present
    for bad_file in ["01_Chapter_1_Introduction.pdf", "02_Regional_Context.pdf", "03_TCAMPO_Organization_and_Structure.pdf"]:
        p = os.path.join(SECTIONS_DIR, bad_file)
        if os.path.exists(p):
            os.remove(p)
            print(f"  Removed outdated section {p}")

def create_full_plan():
    print("2. Generating Full Plan (Chapters 1 & 2) from original document...")
    full_docx_tmp = os.path.join(TMP_DIR, "TCAMPO_Plan2050_Full_Document_Ch1_Ch2.docx")
    shutil.copy(ORIGINAL_DOCX, full_docx_tmp)
    shutil.copy(ORIGINAL_DOCX, os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Full_Document_Ch1_Ch2.docx"))
    
    output_pdf = os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Full_Document_Ch1_Ch2.pdf")
    docx_to_pdf_via_ql(full_docx_tmp, output_pdf)

def create_chapter_1():
    print("3. Generating Chapter 1: Introduction from original document...")
    ch1_docx_tmp = os.path.join(TMP_DIR, "TCAMPO_Plan2050_Chapter1_Introduction.docx")
    shutil.copy(ORIGINAL_DOCX, ch1_docx_tmp)
    doc = docx.Document(ch1_docx_tmp)
    
    found_ch2 = False
    to_remove = []
    for child in list(doc.element.body):
        if child.tag.endswith('p'):
            p = Paragraph(child, doc)
            if p.style.name.startswith('Heading') and p.text.strip().startswith('Chapter 2:'):
                found_ch2 = True
        if found_ch2:
            to_remove.append(child)
            
    for child in to_remove:
        doc.element.body.remove(child)
        
    doc.save(ch1_docx_tmp)
    shutil.copy(ch1_docx_tmp, os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Chapter1_Introduction.docx"))
    
    output_pdf = os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Chapter1_Introduction.pdf")
    docx_to_pdf_via_ql(ch1_docx_tmp, output_pdf)

def create_chapter_2():
    print("4. Generating Chapter 2: Trends and Forecasts from original document...")
    ch2_docx_tmp = os.path.join(TMP_DIR, "TCAMPO_Plan2050_Chapter2_National_and_Regional_Trends_and_Forecasts.docx")
    shutil.copy(ORIGINAL_DOCX, ch2_docx_tmp)
    doc = docx.Document(ch2_docx_tmp)
    
    to_remove = []
    for child in list(doc.element.body):
        if child.tag.endswith('p'):
            p = Paragraph(child, doc)
            if p.style.name.startswith('Heading') and p.text.strip().startswith('Chapter 2:'):
                break
        to_remove.append(child)
        
    for child in to_remove:
        doc.element.body.remove(child)
        
    doc.save(ch2_docx_tmp)
    shutil.copy(ch2_docx_tmp, os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Chapter2_National_and_Regional_Trends_and_Forecasts.docx"))
    
    output_pdf = os.path.join(PUBLIC_DOWNLOADS, "TCAMPO_Plan2050_Chapter2_National_and_Regional_Trends_and_Forecasts.pdf")
    docx_to_pdf_via_ql(ch2_docx_tmp, output_pdf)

def create_sections():
    print("5. Generating all individual section PDFs from original document...")
    sections_map = [
        # Chapter 1 (Parent Header index 0: Heading 2: Chapter 1: Introduction)
        ("TCAMPO_Plan2050_Sec1_1_Regional_Context", 0, 1, 7),
        ("TCAMPO_Plan2050_Sec1_2_TCAMPO_Organization_and_Structure", 0, 7, 11),
        ("TCAMPO_Plan2050_Sec1_3_Purpose_of_the_Plan", 0, 11, 13),
        ("TCAMPO_Plan2050_Sec1_4_Role_of_TCAMPO", 0, 13, 20),
        ("TCAMPO_Plan2050_Sec1_5_Federal_Requirements", 0, 20, 23),
        ("TCAMPO_Plan2050_Sec1_6_National_Planning_Factors", 0, 23, 35),
        ("TCAMPO_Plan2050_Sec1_7_Federal_Performance_Measures", 0, 35, 45),
        ("TCAMPO_Plan2050_Sec1_8_Goals_Objectives_Performance_Measures", 0, 45, 77),
        ("TCAMPO_Plan2050_Sec1_9_Planning_Process", 0, 77, 91),
        ("TCAMPO_Plan2050_Sec1_10_Relationship_to_Other_Plans", 0, 91, 93),
        ("TCAMPO_Plan2050_Sec1_11_What_We_Heard_Public_Outreach", 0, 93, 96),

        # Chapter 2 (Parent Header index 96: Heading 2: Chapter 2: National and Regional Trends and Forecasts)
        ("TCAMPO_Plan2050_Sec2_1_Trends_Introduction", None, 96, 101),
        ("TCAMPO_Plan2050_Sec2_2_Activity_Density", 96, 101, 118),
        ("TCAMPO_Plan2050_Sec2_3_Population_Employment_Trends", 96, 118, 150),
        ("TCAMPO_Plan2050_Sec2_4_Shift_Share_Employment", 96, 150, 161),
        ("TCAMPO_Plan2050_Sec2_5_Demographic_Profiles_EJ", 96, 161, 166),
        ("TCAMPO_Plan2050_Sec2_6_Socioeconomic_Forecasts", 96, 166, 172),
        ("TCAMPO_Plan2050_Sec2_7_Land_Use_and_Development", 96, 172, 184),
        ("TCAMPO_Plan2050_Sec2_8_Travel_Patterns", 96, 184, 196),
        ("TCAMPO_Plan2050_Sec2_9_Commuter_Flows", 96, 196, 212),
    ]

    for filename, parent_idx, start_idx, end_idx in sections_map:
        sec_docx = os.path.join(TMP_DIR, f"{filename}.docx")
        sec_pdf = os.path.join(SECTIONS_DIR, f"{filename}.pdf")
        
        # Only regenerate if missing or size is suspicious
        if os.path.exists(sec_pdf) and os.path.getsize(sec_pdf) > 10000:
            print(f"  [Exists] {filename}.pdf ({os.path.getsize(sec_pdf):,} bytes)")
            continue

        doc = docx.Document(ORIGINAL_DOCX)
        body = doc.element.body
        sectPr = body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
        
        elements = []
        if parent_idx is not None:
            elements.append(body[parent_idx])
        for i in range(start_idx, end_idx):
            elements.append(body[i])
            
        for child in list(body):
            if child != sectPr:
                body.remove(child)
                
        for el in elements:
            if sectPr is not None:
                sectPr.addprevious(el)
            else:
                body.append(el)
                
        doc.save(sec_docx)
        shutil.copy(sec_docx, os.path.join(SECTIONS_DIR, f"{filename}.docx"))
        docx_to_pdf_via_ql(sec_docx, sec_pdf)

def main():
    cleanup_old_ai_docs()
    create_full_plan()
    create_chapter_1()
    create_chapter_2()
    create_sections()
    print("\nAll original document PDFs have been successfully created from the client's Word document!")

if __name__ == "__main__":
    main()
